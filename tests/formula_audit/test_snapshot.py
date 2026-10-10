"""Parity between the saved-workbook (openpyxl) and in-memory (future live-Excel) sources.

The whole point: the *same* detector produces *identical* findings, coverage, ordering and
confidence regardless of where the formulas came from. No second detection pathway.
"""

import pytest
from openpyxl import Workbook, load_workbook

from excel_toolkit.formula_audit import (
    InMemoryWorksheet,
    scan_worksheet,
    snapshot_from_com,
)

from .fixtures_ws import SHEET_LABELS, build_ws_fixtures, expected_cells


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    path = build_ws_fixtures(tmp_path_factory.mktemp("ws") / "corpus.xlsx")
    return load_workbook(path)


def _assert_reports_equal(a, b):
    assert a.sheet == b.sheet
    assert a.findings == b.findings                      # coords, kind, axes, expected, conflict, message
    assert (a.total_formulas, a.resolved, a.unsupported, a.analysed_formulas) == \
           (b.total_formulas, b.resolved, b.unsupported, b.analysed_formulas)
    assert a.unsupported_cells == b.unsupported_cells
    assert a.skipped_cells == b.skipped_cells


@pytest.mark.parametrize("sheet", list(SHEET_LABELS))
def test_saved_and_snapshot_produce_identical_reports(wb, sheet):
    ws = wb[sheet]
    snapshot = InMemoryWorksheet.from_openpyxl(ws)
    saved = scan_worksheet(ws)
    snap = scan_worksheet(snapshot)
    _assert_reports_equal(saved, snap)
    # and the expected anomalies are still found via the snapshot
    assert expected_cells(SHEET_LABELS[sheet]) <= {f.coord for f in snap.findings}


def test_snapshot_carries_sheet_identity():
    snap = InMemoryWorksheet({"B2": "=B1*2"}, title="Budget 2026")
    assert scan_worksheet(snap).sheet == "Budget 2026"


def test_in_memory_snapshot_hand_built_with_constructs():
    # horizontal + vertical anomalies, blanks, constants, actual->forecast, unsupported, identity
    cells = {
        "A1": 0.05,
        # actuals (constants) then forecast (formulas) in row 2 -> boundary, no FP
        "A2": "Rev", "B2": 100, "C2": 120, "D2": "=C2+$A$1", "E2": "=D2+$A$1", "F2": "=E2+$A$1",
        # base row 3 (constants)
        "B3": 10, "C3": 10, "D3": 10, "E3": 10, "F3": 10,
        # row 4 consistent (=X2-X3) with a singleton anomaly at E4 (=E2-D3)
        "A4": "GP", "B4": "=B2-B3", "C4": "=C2-C3", "D4": "=D2-D3", "E4": "=E2-D3", "F4": "=F2-F3",
        # vertical column H consistent (=above+1) with an INDIRECT (unsupported) in the middle
        "H1": 1, "H2": "=H1+1", "H3": '=INDIRECT("A1")', "H4": "=H3+1", "H5": "=H4+1",
    }
    snap = InMemoryWorksheet(cells, title="S")
    report = scan_worksheet(snap)
    flagged = {f.coord for f in report.findings}
    assert "E4" in flagged                               # row singleton found
    assert any(c == "H3" for c, _s in report.unsupported_cells)  # INDIRECT surfaced
    assert report.sheet == "S"


def test_unsaved_edit_is_detected_via_snapshot():
    # The saved file is clean; an unsaved edit introduces an anomaly the snapshot carries.
    saved = {f"{c}1": 1 for c in "BCDE"}
    saved.update({f"{c}2": f"={c}1*2" for c in "BCDE"})
    assert scan_worksheet(InMemoryWorksheet(saved, "S")).findings == []
    unsaved = dict(saved)
    unsaved["D2"] = "=D1*3"                               # edited in Excel, not yet saved
    flagged = {f.coord for f in scan_worksheet(InMemoryWorksheet(unsaved, "S")).findings}
    assert flagged == {"D2"}


# --- bulk snapshot-conversion from a FAKE COM worksheet (no real Excel) -------------------

class _FakeUsedRange:
    """UsedRange with a bulk .Formula (scalar | 1-D | 2-D), like pywin32 returns."""
    def __init__(self, row, col, formula):
        self.Row, self.Column, self.Formula = row, col, formula


class _FakeComWorksheet:
    def __init__(self, name, row, col, formula):
        self.Name, self.UsedRange = name, _FakeUsedRange(row, col, formula)


def test_snapshot_from_fake_com_bulk_matches_in_memory():
    # one bulk 2-D .Formula read: row 1 constants, row 2 formulas with a D2 anomaly
    formula = (
        (1, 1, 1, 1),
        ("=B1*2", "=C1*2", "=D1*3", "=E1*2"),
    )
    com = _FakeComWorksheet("Live", 1, 2, formula)      # top-left B1
    reference = InMemoryWorksheet(
        {"B1": 1, "C1": 1, "D1": 1, "E1": 1,
         "B2": "=B1*2", "C2": "=C1*2", "D2": "=D1*3", "E2": "=E1*2"}, "Live")
    a, b = scan_worksheet(snapshot_from_com(com)), scan_worksheet(reference)
    _assert_reports_equal(a, b)
    assert {f.coord for f in a.findings} == {"D2"}
    assert a.sheet == "Live"


def test_snapshot_from_com_single_cell_scalar():
    # a single-cell UsedRange returns a scalar, not a 2-D array
    snap = snapshot_from_com(_FakeComWorksheet("S", 5, 3, "=A1+1"))
    assert snap.max_row == 5 and snap.max_col == 3
    assert snap.value_at(5, 3) == "=A1+1"


def test_snapshot_from_com_handles_constants_blanks_errors_dates():
    import datetime
    err = object()  # stand-in for an Excel error / unexpected COM type
    formula = (
        ("=A1", 42, None),
        ("text", err, datetime.datetime(2026, 1, 31)),
    )
    snap = snapshot_from_com(_FakeComWorksheet("S", 1, 1, formula))
    assert snap.value_at(1, 1) == "=A1"                 # formula kept
    assert snap.value_at(1, 2) == 42                    # number constant
    assert snap.value_at(1, 3) is None                  # blank omitted
    assert snap.value_at(2, 1) == "text"                # text constant
    assert snap.value_at(2, 2) == str(err)              # error -> string (boundary)
    assert snap.value_at(2, 3) == datetime.datetime(2026, 1, 31)  # date kept
