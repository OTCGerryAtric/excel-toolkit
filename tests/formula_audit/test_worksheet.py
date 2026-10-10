"""Worksheet auto-scan: singleton + block detection, axes, dedup/order, coverage, metrics."""

import pytest
from openpyxl import load_workbook

from excel_toolkit.formula_audit import openpyxl_supported, scan_worksheet

from .fixtures_ws import (
    SHEET_LABELS,
    build_ws_fixtures,
    expected_cells,
    known_fn_cells,
)


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    path = build_ws_fixtures(tmp_path_factory.mktemp("ws") / "corpus.xlsx")
    return load_workbook(path)


def _flagged_groups(report):
    """Flagged cells grouped into blocks (singletons are size-1 groups)."""
    groups = set()
    for f in report.findings:
        groups.add(frozenset(f.block) if (f.kind == "block" and f.block) else frozenset({f.coord}))
    return groups


def test_openpyxl_version_constrained():
    assert openpyxl_supported(), "detector relies on openpyxl 3.1.x internals"


@pytest.mark.parametrize("sheet", list(SHEET_LABELS))
def test_flags_expected_and_no_false_positives(wb, sheet):
    label = SHEET_LABELS[sheet]
    report = scan_worksheet(wb[sheet])
    flagged = {f.coord for f in report.findings}
    exp, kfn = expected_cells(label), known_fn_cells(label)
    assert exp <= flagged, f"{sheet}: missed {exp - flagged}"
    assert flagged & kfn == set(), f"{sheet}: flagged a known-FN cell {flagged & kfn}"
    assert flagged - (exp | kfn) == set(), f"{sheet}: false positives {flagged - (exp | kfn)}"


def test_axes_recorded(wb):
    for sheet, label in SHEET_LABELS.items():
        by_coord = {f.coord: f for f in scan_worksheet(wb[sheet]).findings}
        for coord, axes in label["axes"].items():
            assert by_coord[coord].axes == axes, f"{sheet}!{coord}: {by_coord[coord].axes} != {axes}"


def test_block_findings_carry_evidence(wb):
    for sheet, run in (("consecutive_block", {"D2", "E2"}), ("consec_vertical", {"B5", "B6"})):
        report = scan_worksheet(wb[sheet])
        block_cells = {f.coord for f in report.findings if f.kind == "block"}
        assert block_cells == run
        for f in report.findings:
            if f.kind == "block":
                assert set(f.block) == run          # evidence: the whole run preserved
                assert f.expected and f.observed and f.expected != f.observed
                assert f.confidence == "high"


def test_triple_run_is_suppressed(wb):
    report = scan_worksheet(wb["triple_suppressed"])
    assert report.findings == []  # 3-cell run exceeds the short-block max -> conservative FN


@pytest.mark.parametrize("sheet", ["quarterly", "fy_boundary", "actual_forecast"])
def test_legitimate_multiperiod_variations_not_flagged(wb, sheet):
    assert scan_worksheet(wb[sheet]).findings == []


def test_findings_deterministic_order_and_deduped(wb):
    report = scan_worksheet(wb["overlap"])
    coords = [f.coord for f in report.findings]
    assert coords == sorted(coords, key=lambda c: (int(c[1:]), c[0]))
    assert len(coords) == len(set(coords))


def test_findings_are_potential_not_certain(wb):
    report = scan_worksheet(wb["revenue_grid"])
    assert report.findings
    for f in report.findings:
        assert "Potential inconsistency" in f.message and "error" not in f.message.lower()
        assert f.confidence == "high"


def test_coverage_axes_reported_separately(wb):
    report = scan_worksheet(wb["revenue_grid"])
    assert 0.0 <= report.pattern_coverage <= report.dependency_coverage <= 1.0
    assert report.dependency_coverage == 1.0
    assert report.total_formulas > 0


def test_totals_headings_no_false_positives(wb):
    assert scan_worksheet(wb["totals_headings"]).findings == []


def test_isolated_formula_is_skipped_not_flagged():
    from openpyxl import Workbook
    ws = Workbook().active
    ws["C3"] = "=A3+1"
    report = scan_worksheet(ws)
    assert report.findings == []
    assert report.total_formulas == 1
    assert report.analysed_formulas == 0
    assert ("C3", "insufficient series context") in report.skipped_cells


def test_scanning_does_not_materialise_empty_cells():
    from openpyxl import Workbook
    ws = Workbook().active
    for r in range(1, 5):
        ws[f"B{r}"] = r
    ws["C2"], ws["C3"], ws["C4"] = "=B2+1", "=B3+1", "=B4+1"
    ws["Z500"] = 1
    before = len(ws._cells)
    report = scan_worksheet(ws)
    assert len(ws._cells) == before, "scan created empty cells"
    assert report.total_formulas == 3


def test_unsupported_formula_in_series_not_counted_as_analysed():
    from openpyxl import Workbook
    ws = Workbook().active
    ws["B2"], ws["C2"] = "=B1*2", "=C1*2"
    ws["D2"] = '=INDIRECT("A1")*2'
    ws["E2"] = "=E1*2"
    report = scan_worksheet(ws)
    assert report.total_formulas == 4
    assert report.resolved == 3
    assert report.analysed_formulas == 3
    assert report.pattern_coverage == 3 / 4
    assert ("D2", "unsupported") in report.unsupported_cells
    assert report.findings == []


def test_cross_axis_conflict_is_reported_not_silently_resolved():
    from openpyxl import Workbook
    ws = Workbook().active
    for col in ("B", "C", "E"):
        ws[f"{col}1"] = 10
    for r in (2, 4, 5):
        ws[f"A{r}"] = r
    ws["D1"] = 1
    ws["B3"], ws["C3"], ws["E3"] = "=B$1*2", "=C$1*2", "=E$1*2"
    ws["D2"], ws["D4"], ws["D5"] = "=$A2*3", "=$A4*3", "=$A5*3"
    ws["D3"] = "=D1+D2"
    d3 = next(f for f in scan_worksheet(ws).findings if f.coord == "D3")
    assert set(d3.axes) == {"row", "column"}
    assert d3.conflict is True
    by_axis = dict(d3.expected_by_axis)
    assert by_axis["row"] != by_axis["column"]
    assert "disagree" in d3.message


def test_metrics_cell_and_block_level(wb, capsys):
    tp = fp = fn = 0
    tpb = fpb = fnb = 0
    resolved = total = analysed = 0
    for sheet, label in SHEET_LABELS.items():
        report = scan_worksheet(wb[sheet])
        flagged = {f.coord for f in report.findings}
        exp, kfn = expected_cells(label), known_fn_cells(label)
        tp += len(flagged & exp)
        fp += len(flagged - (exp | kfn))
        fn += len(exp - flagged) + len(kfn)
        resolved += report.resolved
        total += report.total_formulas
        analysed += report.analysed_formulas

        groups = _flagged_groups(report)
        exp_blocks = [frozenset(b) for b in label["expected_blocks"]]
        kfn_blocks = [frozenset(b) for b in label["known_fn_blocks"]]
        tpb += sum(1 for b in exp_blocks if b in groups)
        fnb += sum(1 for b in exp_blocks if b not in groups) + len(kfn_blocks)
        labelled = set(exp_blocks) | set(kfn_blocks)
        fpb += sum(1 for g in groups if g not in labelled)

    prec_c = tp / (tp + fp) if (tp + fp) else 1.0
    rec_c = tp / (tp + fn) if (tp + fn) else 1.0
    prec_b = tpb / (tpb + fpb) if (tpb + fpb) else 1.0
    rec_b = tpb / (tpb + fnb) if (tpb + fnb) else 1.0
    with capsys.disabled():
        print(f"\n  corpus formula cells  : {total}")
        print(f"  CELL : TP={tp} FP={fp} FN={fn}  precision={prec_c:.0%} recall={rec_c:.0%}")
        print(f"  BLOCK: TP={tpb} FP={fpb} FN={fnb}  precision={prec_b:.0%} recall={rec_b:.0%}")
        print(f"         (FN = the deliberately-suppressed 3-cell run)")
        print(f"  dependency coverage   : {resolved / total:.0%}")
        print(f"  pattern coverage      : {analysed / total:.0%}")

    assert fp == 0 and fpb == 0            # no false positives at either level
    assert prec_c == 1.0 and prec_b == 1.0
    assert total > 26
