"""Build offline .xlsx workbook fixtures for the formula-inconsistency detector.

Each sheet is one scenario. Genuine reference mistakes sit alongside *legitimate*
variations that must NOT be flagged: period actuals→forecast boundaries, year-end
totals, and edge cells. Formulas are written as strings; nothing is evaluated and no
Excel instance is involved.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook


def build_fixtures(path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    # 1. clean: four months, one consistent growth formula. No mistakes.
    ws = wb.create_sheet("clean")
    for col in "BCDE":
        ws[f"{col}2"] = 100
        ws[f"{col}3"] = 0.05
        ws[f"{col}4"] = f"={col}2*(1+{col}3)"

    # 2. march: March references row 4 instead of row 3 — a genuine reference mistake.
    ws = wb.create_sheet("march")
    for col in "BCDE":
        ws[f"{col}2"] = 100
        ws[f"{col}3"] = 0.05
        ws[f"{col}4"] = f"={col}2*(1+{col}3)"
    ws["D4"] = "=D2*(1+D4)"  # the deliberate error (D4 should be D3)

    # 3. actuals -> forecast: first two months are hard-coded actuals (legitimate),
    #    later months are a consistent forecast formula. The boundary must NOT flag.
    ws = wb.create_sheet("actuals")
    ws["A1"] = 0.05  # growth rate
    ws["B4"] = 100   # actual
    ws["C4"] = 120   # actual
    ws["D4"] = "=C4*(1+$A$1)"
    ws["E4"] = "=D4*(1+$A$1)"
    ws["F4"] = "=E4*(1+$A$1)"

    # 4. year-end total: consistent months then a SUM total at the end. Edge, must NOT flag.
    ws = wb.create_sheet("total")
    for col in "BCDE":
        ws[f"{col}2"] = 100
        ws[f"{col}3"] = 0.05
        ws[f"{col}4"] = f"={col}2*(1+{col}3)"
    ws["F4"] = "=SUM(B4:E4)"

    # 5. absolute_break: all months use the absolute rate $B$1, except March which points
    #    at $C$1 by mistake. Tests that absolute refs are compared correctly.
    ws = wb.create_sheet("absolute")
    ws["B1"] = 0.05
    ws["C1"] = 0.08
    for col in "BCDE":
        ws[f"{col}3"] = 100
        ws[f"{col}4"] = f"={col}3*(1+$B$1)"
    ws["D4"] = "=D3*(1+$C$1)"  # wrong absolute cell

    # 6. indirect: one cell uses INDIRECT (run-time dynamic) -> unsupported, not flagged,
    #    not counted clean.
    ws = wb.create_sheet("indirect")
    for col in "BCDE":
        ws[f"{col}2"] = 100
        ws[f"{col}3"] = 0.05
        ws[f"{col}4"] = f"={col}2*(1+{col}3)"
    ws["C4"] = '=INDIRECT("B2")*(1+C3)'

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# scenario -> (sheet, row, (start_col, end_col), genuine mistakes, legitimate cells, unsupported)
SCENARIOS = {
    "clean":    ("clean",    4, (2, 5), set(),       {"B4", "C4", "D4", "E4"},       set()),
    "march":    ("march",    4, (2, 5), {"D4"},      {"B4", "C4", "E4"},             set()),
    "actuals":  ("actuals",  4, (2, 6), set(),       {"B4", "C4", "D4", "E4", "F4"}, set()),
    "total":    ("total",    4, (2, 6), set(),       {"B4", "C4", "D4", "E4", "F4"}, set()),
    "absolute": ("absolute", 4, (2, 5), {"D4"},      {"B4", "C4", "E4"},             set()),
    "indirect": ("indirect", 4, (2, 5), set(),       {"B4", "D4", "E4"},             {"C4"}),
}
