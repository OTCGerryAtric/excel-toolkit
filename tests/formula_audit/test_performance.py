"""Performance/memory against a larger offline workbook, and source-unchanged guarantee."""

import hashlib
import time
import tracemalloc

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import get_column_letter

from excel_toolkit.formula_audit import scan_worksheet

ROWS, COLS = 200, 100  # 200 * 99 = 19,800 formula cells


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _build_large(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "big"
    for r in range(1, ROWS + 1):
        ws[f"A{r}"] = r  # numeric seed column
        for c in range(2, COLS + 1):
            ws[f"{get_column_letter(c)}{r}"] = f"={get_column_letter(c - 1)}{r}+1"
    # three interior anomalies (reference two columns left instead of one)
    planted = set()
    for (r, c) in [(100, 50), (37, 80), (160, 25)]:
        coord = f"{get_column_letter(c)}{r}"
        ws[coord] = f"={get_column_letter(c - 2)}{r}+1"
        planted.add(coord)
    wb.save(path)
    return planted


def test_performance_memory_and_source_unchanged(tmp_path, capsys):
    path = tmp_path / "big.xlsx"
    planted = _build_large(path)

    before = _sha256(path)
    wb = load_workbook(path)

    tracemalloc.start()
    start = time.perf_counter()
    report = scan_worksheet(wb["big"])
    elapsed = time.perf_counter() - start
    _current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    after = _sha256(path)  # we only read; the file must be byte-identical
    flagged = {f.coord for f in report.findings}

    with capsys.disabled():
        print(f"\n  formula cells scanned : {report.total_formulas}")
        print(f"  scan time             : {elapsed:.2f} s")
        print(f"  peak memory (scan)    : {peak / 1e6:.1f} MB")
        print(f"  findings              : {len(flagged)} (planted {len(planted)})")

    assert before == after                         # source workbook unchanged
    assert report.total_formulas == ROWS * (COLS - 1)
    assert flagged == planted                      # found exactly the planted anomalies, no FP
    assert elapsed < 60                            # generous ceiling for CI
