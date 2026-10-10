"""Worksheet-level regression corpus: realistic financial-model layouts with labelled
anomalies and legitimate variations (headings, totals, actuals, opening balances).

Each sheet declares the cells that SHOULD be flagged (expected_tp), anomalies we expect
to MISS in this increment (known_fn, e.g. consecutive identical errors), and the axes a
finding should carry. Everything else is a legitimate cell that must NOT be flagged.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter


def _revenue_grid(ws):
    # Assumptions
    ws["A20"], ws["B20"] = "Growth", 0.05
    ws["A21"], ws["B21"] = "COGS %", 0.40
    # Header row (text) + row labels (text) -> structural boundaries
    ws["A1"] = "Item"
    for i, month in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
        ws[f"{get_column_letter(2 + i)}1"] = month
    ws["H1"] = "Total"
    # Revenue: Jan actual (constant) then consistent growth; year total.
    ws["A2"], ws["B2"] = "Revenue", 1000
    for c in range(3, 8):  # C..G
        prev = get_column_letter(c - 1)
        ws[f"{get_column_letter(c)}2"] = f"={prev}2*(1+$B$20)"
    ws["H2"] = "=SUM(B2:G2)"
    # COGS as % of revenue.
    ws["A3"], ws["B3"] = "COGS", 400
    for c in range(3, 8):
        ws[f"{get_column_letter(c)}3"] = f"={get_column_letter(c)}2*$B$21"
    ws["H3"] = "=SUM(B3:G3)"
    # Gross profit = revenue - COGS, with a planted error in Apr (E4 -> D3 not E3).
    ws["A4"] = "Gross Profit"
    for c in range(2, 8):  # B..G
        col = get_column_letter(c)
        ws[f"{col}4"] = f"={col}2-{col}3"
    ws["E4"] = "=E2-D3"          # <-- anomaly
    ws["H4"] = "=SUM(B4:G4)"


def _balance_roll(ws):
    ws["D1"] = 10  # increment
    ws["A2"], ws["B2"] = "Opening", 100
    for r in range(3, 9):  # B3..B8 roll forward
        ws[f"A{r}"] = f"Period {r - 2}"
        ws[f"B{r}"] = f"=B{r - 1}+$D$1"
    ws["B6"] = "=B5*$D$1"  # <-- anomaly (multiplies instead of adds)


def _overlap(ws):
    # 2-D fill: cell = rowfactor * colfactor; one cell wrong in both directions.
    for c in range(2, 6):  # B..E col headers
        ws[f"{get_column_letter(c)}1"] = c * 10
    for r in range(2, 6):  # A2..A5 row factors
        ws[f"A{r}"] = r
    for r in range(2, 6):
        for c in range(2, 6):
            col = get_column_letter(c)
            ws[f"{col}{r}"] = f"=$A{r}*{col}$1"
    ws["D3"] = "=$A3*C$1"  # <-- anomaly (C$1 should be D$1), interior in row 3 AND column D


def _consecutive_fn(ws):
    for c in range(2, 8):  # B1..G1 headers
        ws[f"{get_column_letter(c)}1"] = c
    ws["A2"] = "Doubler"
    for c in range(2, 8):  # B2..G2 = header*2
        ws[f"{get_column_letter(c)}2"] = f"={get_column_letter(c)}1*2"
    ws["D2"] = "=D1*3"  # two adjacent identical errors -> neither is a flanked singleton
    ws["E2"] = "=E1*3"


def _totals_headings(ws):
    # drivers
    for c in range(2, 5):
        ws[f"{get_column_letter(c)}10"] = 5
        ws[f"{get_column_letter(c)}11"] = 3
    ws["A1"] = "Metric"
    for i, m in enumerate(["M1", "M2", "M3"]):
        ws[f"{get_column_letter(2 + i)}1"] = m
    ws["E1"] = "Total"
    ws["A2"] = "Sales"
    for c in range(2, 5):
        ws[f"{get_column_letter(c)}2"] = f"={get_column_letter(c)}10*2"
    ws["E2"] = "=SUM(B2:D2)"
    ws["A3"] = "Costs"
    for c in range(2, 5):
        ws[f"{get_column_letter(c)}3"] = f"={get_column_letter(c)}11*2"
    ws["E3"] = "=SUM(B3:D3)"
    # row 4 left blank (structural boundary)
    ws["A5"] = "Total"
    for c in range(2, 6):
        col = get_column_letter(c)
        ws[f"{col}5"] = f"=SUM({col}2:{col}3)"


def build_ws_fixtures(path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, builder in (
        ("revenue_grid", _revenue_grid),
        ("balance_roll", _balance_roll),
        ("overlap", _overlap),
        ("consecutive_fn", _consecutive_fn),
        ("totals_headings", _totals_headings),
    ):
        builder(wb.create_sheet(name))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# sheet -> (expected true positives, known/expected false negatives, {coord: expected axes})
SHEET_LABELS = {
    "revenue_grid":   ({"E4"},        set(),          {"E4": ("row",)}),
    "balance_roll":   ({"B6"},        set(),          {"B6": ("column",)}),
    "overlap":        ({"D3"},        set(),          {"D3": ("column", "row")}),
    "consecutive_fn": (set(),         {"D2", "E2"},   {}),
    "totals_headings": (set(),        set(),          {}),
}
