"""M3d manual investigation: write the disposable test workbook the owner opens in normal Excel.

Written with openpyxl - no Excel is started. The workbook is a plain .xlsx with no macros; the owner imports
the probe (m3d_paste_probe.bas) into it by hand for the session and closes it without saving.

Sheets (checklist in docs/validation/m3d-paste-values-experiment.md, "Manual investigation"):
- "1 Baseline", "2 Probe paste", "3 Do nothing": source A1:B3 as in the automated run, destination D1:E3
  "old", history cells J1 (and J2 on "3 Do nothing").
- "4a Filter macro", "4b Filter Excel": the automated filtered layout - A1:B3 source, D1:H6 "old", D1:D6
  head/a/b/a/b/a filtered to "a" (rows 3 and 5 hidden, which hides source cell A3 too). Row 8 reads E1:E6
  back and row 11 reads A1:A3 back, so hidden cells can be seen without clearing the filter.
- "4c Filter other source" with its source on "4c Source" (not filtered): separates the source effect from
  the destination effect.

  <scratch venv>/python addin/live/m3d_manual_book.py <output.xlsx>
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

SOURCE = {"A1": (1234.5, "#,##0.00"), "A2": ("=A1*2", None), "A3": ("text", None),
          "B1": (45000, "dd/mm/yyyy"), "B2": (0.25, "0%"), "B3": (True, None)}
FILTER_COLUMN = ("head", "a", "b", "a", "b", "a")
HIDDEN_ROWS = (3, 5)


def _source(ws) -> None:
    for address, (content, number_format) in SOURCE.items():
        ws[address] = content
        if number_format:
            ws[address].number_format = number_format


def _undo_sheet(wb, name: str, history_cells=("J1",)) -> None:
    ws = wb.create_sheet(name)
    _source(ws)
    for row in ws["D1:E3"]:
        for cell in row:
            cell.value = "old"
    for address in history_cells:
        ws[address] = "history"
    ws["A5"] = "Source A1:B3. Destination D1. History cell(s): " + ", ".join(history_cells)


def _filter_sheet(wb, name: str, source_sheet: str | None) -> None:
    ws = wb.create_sheet(name)
    if source_sheet is None:
        _source(ws)
    for row in ws["D1:H6"]:
        for cell in row:
            cell.value = "old"
    for row, value in enumerate(FILTER_COLUMN, start=1):
        ws.cell(row=row, column=4, value=value)
    ws.auto_filter.ref = "D1:D6"
    ws.auto_filter.add_filter_column(0, ["a"])
    for row in HIDDEN_ROWS:
        ws.row_dimensions[row].hidden = True
    ws["A8"] = "E1:E6 now:"
    ws["A8"].font = Font(bold=True)
    for offset in range(6):
        ws.cell(row=8, column=2 + offset, value=f"=E{offset + 1}")
        ws.cell(row=7, column=2 + offset, value=f"E{offset + 1}" + (" (hidden)" if offset + 1 in HIDDEN_ROWS else ""))
    origin = f"'{source_sheet}'!" if source_sheet else ""
    ws["A11"] = f"Source {origin or ''}A1:A3 now:"
    ws["A11"].font = Font(bold=True)
    for offset in range(3):
        ws.cell(row=11, column=2 + offset, value=f"={origin}A{offset + 1}")
        ws.cell(row=10, column=2 + offset, value=f"A{offset + 1}" + (" (hidden)" if source_sheet is None and offset + 1 in HIDDEN_ROWS else ""))


def build(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _undo_sheet(wb, "1 Baseline")
    _undo_sheet(wb, "3 Do nothing", history_cells=("J1", "J2"))
    _undo_sheet(wb, "2 Probe paste")
    _filter_sheet(wb, "4a Filter macro", None)
    _filter_sheet(wb, "4b Filter Excel", None)
    src = wb.create_sheet("4c Source")
    _source(src)
    _filter_sheet(wb, "4c Filter other source", "4c Source")
    wb.save(path)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: m3d_manual_book.py <output.xlsx>")
    build(Path(sys.argv[1]))
    print("written:", sys.argv[1])
