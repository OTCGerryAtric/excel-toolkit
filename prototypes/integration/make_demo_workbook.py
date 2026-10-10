"""Generate a disposable synthetic workbook for the Stage 2 live-Excel demo.

Offline only (openpyxl, no Excel). No business information. Contains a horizontal series,
a vertical series, a *planted* saved anomaly, and a clean cell the tester edits (unsaved)
to introduce a second anomaly during the live test.

    uv run python prototypes/integration/make_demo_workbook.py [out.xlsx]
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

SHEET = "Demo"
EDIT_CELL = "D3"
EDIT_FROM = "=C3*(1+$A$10)"
EDIT_TO = "=C3*(1+$A$11)"   # wrong rate cell ($A$11 instead of $A$10) -> unsaved anomaly
PLANTED_CELL = "D4"          # saved anomaly (references C3 instead of D3)


def build(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    ws["A9"], ws["A10"] = "Growth", 0.05           # assumption (A10 = growth rate)
    ws["A1"] = "Line"
    for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun"]):
        ws[f"{chr(66 + i)}1"] = m                   # B1..G1
    ws["I1"] = "H1 Total"

    # Revenue (row 3): clean growth series Jan..Jun; the tester edits D3 (unsaved).
    ws["A3"], ws["B3"] = "Revenue", 1000
    for c in range(3, 8):                            # C3..G3
        ws[f"{chr(64 + c)}3"] = f"={chr(63 + c)}3*(1+$A$10)"
    ws["I3"] = "=SUM(B3:G3)"                         # total (column H left empty as a gap)

    # Margin (row 4): consistent =X3*0.3 with a PLANTED anomaly at D4 (=C3*0.3).
    ws["A4"] = "Margin"
    for c in range(2, 8):                            # B4..G4
        ws[f"{chr(64 + c)}4"] = f"={chr(64 + c)}3*0.3"
    ws["D4"] = "=C3*0.3"                             # planted: should be =D3*0.3
    ws["I4"] = "=SUM(B4:G4)"

    # Vertical roll-forward (column K): clean, for orientation coverage.
    ws["K1"], ws["K2"] = "Roll", 100
    for r in range(3, 7):
        ws[f"K{r}"] = f"=K{r - 1}+$A$10"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("demo_workbook.xlsx")
    print(build(out))
