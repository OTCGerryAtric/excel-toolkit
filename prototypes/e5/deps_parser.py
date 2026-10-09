"""E5 feasibility prototype: static dependency extraction from a formula string.

Not the production parser. Uses openpyxl's tokenizer (MIT) and workbook metadata
(defined names, tables, sheet order) - no calculation, no Excel.

Every result carries a status:
  resolved      - every dependency identified
  partial       - some dependencies identified; others depend on run-time values
  indeterminate - the dependencies cannot be established statically
Never "no dependencies" because something wasn't recognised.
"""

import re
from dataclasses import dataclass, field

from openpyxl.formula.tokenizer import Tokenizer

CELL = r"\$?[A-Z]{1,3}\$?\d+"
A1_RANGE = re.compile(rf"^({CELL})(?::({CELL}))?$")
COLUMNS = re.compile(r"^\$?([A-Z]{1,3}):\$?([A-Z]{1,3})$")
ROWS = re.compile(r"^\$?(\d+):\$?(\d+)$")
SPILL = re.compile(rf"((?:'[^']+'|[A-Za-z_][\w. ]*)?!?{CELL})#")
TABLE = re.compile(r"^([A-Za-z_\\][\w.]*)\[(.*)\]$")
NAME = re.compile(r"^[A-Za-z_\\][\w.]*$")
STATUS_ORDER = {"resolved": 0, "partial": 1, "indeterminate": 2}


@dataclass
class Context:
    sheet: str                      # sheet holding the formula
    sheet_order: list[str]
    names: dict[str, str]           # name -> RefersTo, e.g. "=Inputs!$B$1"
    tables: dict[str, dict]         # table -> {"sheet": .., "columns": {column: "B2:B6"}}


@dataclass
class Result:
    dependencies: set[str] = field(default_factory=set)
    status: str = "resolved"
    notes: list[str] = field(default_factory=list)
    absolute: dict[str, str] = field(default_factory=dict)  # reference -> "abs"/"rel"/"mixed"

    def downgrade(self, status: str, note: str) -> None:
        if STATUS_ORDER[status] > STATUS_ORDER[self.status]:
            self.status = status
        self.notes.append(note)


def _strip(cell: str) -> str:
    return cell.replace("$", "")


def _split_sheet(operand: str) -> tuple[str | None, str | None, str]:
    """-> (external book or None, sheet or None, address)"""
    if "!" not in operand:
        return None, None, operand
    prefix, address = operand.rsplit("!", 1)
    prefix = prefix[1:-1].replace("''", "'") if prefix.startswith("'") else prefix
    external = re.match(r"^(?:.*\\)?\[(.+?)\](.+)$", prefix)
    if external:
        return external.group(1), external.group(2), address
    return None, prefix, address


def _add_address(result: Result, book: str | None, sheet: str, address: str, original: str) -> bool:
    plain = _strip(address)
    prefix = f"[{book}]{sheet}" if book else sheet
    if A1_RANGE.match(address) or COLUMNS.match(address) or ROWS.match(address):
        result.dependencies.add(f"{prefix}!{plain}")
        if A1_RANGE.match(address):
            result.absolute[f"{prefix}!{plain}"] = (
                "abs" if address.count("$") == 2 * (1 + (":" in address))
                else "rel" if "$" not in address else "mixed")
        return True
    return False


def dependencies(formula: str, context: Context) -> Result:
    result = Result()
    if not formula.startswith("="):
        return result

    # The tokenizer can't read the spill operator: note spill anchors, then remove '#'.
    spill_anchors = set()
    for match in SPILL.finditer(re.sub(r'"[^"]*"', "", formula)):
        spill_anchors.add(match.group(1))
    text = SPILL.sub(r"\1", formula)

    try:
        tokens = Tokenizer(text).items
    except Exception as error:  # noqa: BLE001
        result.downgrade("indeterminate", f"formula could not be parsed: {error}")
        return result

    # LET: names defined in odd argument positions are local, not workbook names.
    let_locals: set[str] = set()
    stack: list[list] = []  # [function name, argument index]
    for index, token in enumerate(tokens):
        if token.type == "FUNC" and token.subtype == "OPEN":
            stack.append([token.value[:-1].upper(), 0])
        elif token.type == "FUNC" and token.subtype == "CLOSE":
            stack.pop()
        elif token.type == "SEP" and token.subtype == "ARG" and stack:
            stack[-1][1] += 1
        elif token.type == "OPERAND" and token.subtype == "RANGE" and stack and stack[-1][0] == "LET":
            next_token = tokens[index + 1] if index + 1 < len(tokens) else None
            if stack[-1][1] % 2 == 0 and next_token is not None and next_token.type == "SEP":
                let_locals.add(token.value.upper())

    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            if name == "INDIRECT":
                result.downgrade("indeterminate", "INDIRECT: target is text built at run time")
            elif name == "OFFSET":
                result.downgrade("partial", "OFFSET: anchor known, actual target depends on its arguments at run time")
            continue
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        operand = token.value
        if operand.upper() in let_locals:
            continue
        book, sheet, address = _split_sheet(operand)
        sheet = sheet or context.sheet

        if operand in spill_anchors or f"{sheet}!{operand}" in spill_anchors:
            result.dependencies.add(f"{sheet}!{_strip(address)}#")
            result.downgrade("partial", f"spill reference {operand}#: anchor known, extent only known at run time")
            continue
        if sheet and ":" in sheet and not book:  # 3D reference
            first, last = sheet.split(":", 1)
            try:
                span = context.sheet_order[context.sheet_order.index(first): context.sheet_order.index(last) + 1]
            except ValueError:
                result.downgrade("indeterminate", f"3D reference {operand}: sheet range not found")
                continue
            for each in span:
                _add_address(result, None, each, address, operand)
            continue
        if _add_address(result, book, sheet, address, operand):
            continue
        table = TABLE.match(operand)
        if table and table.group(1) in context.tables:
            meta = context.tables[table.group(1)]
            column = table.group(2)
            if column in meta["columns"]:
                result.dependencies.add(f"{meta['sheet']}!{meta['columns'][column]}")
            else:
                result.dependencies.add(f"{meta['sheet']}!<table {table.group(1)}>")
                result.downgrade("partial", f"structured reference {operand}: specifier not resolved by the prototype")
            continue
        if NAME.match(operand) and operand in context.names:
            refers_to = context.names[operand]
            inner = dependencies(refers_to, Context(context.sheet, context.sheet_order, {}, context.tables))
            result.dependencies |= inner.dependencies
            if inner.status != "resolved" or not inner.dependencies:
                result.downgrade("partial", f"name {operand} refers to {refers_to}, not fully resolved")
            continue
        result.downgrade("indeterminate", f"unrecognised reference {operand!r}")
    return result
