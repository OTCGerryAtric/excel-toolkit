"""Automatic worksheet scanning: discover formula series across rows and columns,
segment them conservatively, and apply the high-confidence flanked-singleton rule.

Series discovery
----------------
A *line* is one worksheet row or column. Each line is split by ``segment_line`` into
maximal runs of adjacent cells of the **same kind**:
- ``formula`` — the cell holds a formula (``=...``);
- ``const``   — the cell holds a plain number.
Blanks, text labels/headings, dates and anything else are **boundaries** that end a run,
and a run also ends where the kind changes (so an actuals→forecast transition, i.e. a
numbers run meeting a formulas run, is split). Only ``formula`` segments are analysed.

Reads are **non-mutating**: values are pulled from the worksheet's existing-cell store
with ``dict.get`` (never ``ws.cell()`` / ``iter_rows``, both of which instantiate empty
``Cell`` objects via ``_get_cell``), so scanning a sparse sheet does not grow the
in-memory workbook.

Coverage is reported on two independent axes, kept separate on purpose:
- **dependency coverage** = formulas whose references normalised / total formulas;
- **pattern coverage**    = formulas that were *successfully comparable* (normalised AND
  sat in a qualifying series) / total formulas. An unsupported formula inside an
  otherwise-valid series is NOT counted as analysed — it is reported as unsupported.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from .detector import dominant_pattern, flanked_outliers
from .normalize import RESOLVED, normalize_cell_value


def _member_kind(value) -> str | None:
    """'formula', 'const' (plain number), or None (boundary: blank/text/date/other)."""
    if isinstance(value, str):
        return "formula" if value.startswith("=") else None
    if isinstance(value, bool):
        return "const"
    if isinstance(value, (int, float)):
        return "const"
    return None


def segment_line(values: list) -> list[tuple[int, int, str]]:
    """Split a line's raw values into homogeneous runs.

    Returns inclusive ``(start, end, kind)`` spans where ``kind`` is 'formula' or 'const'.
    Runs end at any boundary cell and wherever the kind changes.
    """
    segments: list[tuple[int, int, str]] = []
    start: int | None = None
    kind: str | None = None
    for i, value in enumerate(values):
        k = _member_kind(value)
        if k is None:
            if start is not None:
                segments.append((start, i - 1, kind))
                start, kind = None, None
        elif start is None:
            start, kind = i, k
        elif k != kind:
            segments.append((start, i - 1, kind))
            start, kind = i, k
    if start is not None:
        segments.append((start, len(values) - 1, kind))
    return segments


@dataclass(frozen=True)
class WorksheetFinding:
    coord: str
    row: int
    col: int
    axes: tuple[str, ...]                        # ("row",), ("column",) or ("column", "row")
    observed: str                                # the flagged cell's own pattern
    expected: str                                # representative expected pattern (first axis)
    expected_by_axis: tuple[tuple[str, str], ...]  # ((axis, expected), ...) sorted by axis
    conflict: bool                               # row and column disagree on the expected pattern
    confidence: str
    message: str


@dataclass
class WorksheetReport:
    total_formulas: int = 0
    resolved: int = 0
    unsupported: int = 0
    analysed_formulas: int = 0  # comparable: normalised AND in a qualifying series
    findings: list[WorksheetFinding] = field(default_factory=list)
    unsupported_cells: list[tuple[str, str]] = field(default_factory=list)
    skipped_cells: list[tuple[str, str]] = field(default_factory=list)

    @property
    def dependency_coverage(self) -> float:
        return self.resolved / self.total_formulas if self.total_formulas else 0.0

    @property
    def pattern_coverage(self) -> float:
        return self.analysed_formulas / self.total_formulas if self.total_formulas else 0.0


def _lines(ws):
    """Yield ('row'|'column', [(coord, value), ...]) for every row then every column.

    Reads from ``ws._cells`` with ``.get`` so unpopulated positions come back as ``None``
    without instantiating ``Cell`` objects (openpyxl's ``ws.cell()`` and ``iter_rows`` both
    create cells via ``_get_cell``; we must not mutate the workbook while scanning it).
    """
    cells = ws._cells  # {(row, col): Cell}; .get never creates a cell
    max_row, max_col = ws.max_row, ws.max_column

    def value_at(r: int, c: int):
        cell = cells.get((r, c))
        return cell.value if cell is not None else None

    for r in range(1, max_row + 1):
        yield "row", [(f"{get_column_letter(c)}{r}", value_at(r, c))
                      for c in range(1, max_col + 1)]
    for c in range(1, max_col + 1):
        letter = get_column_letter(c)
        yield "column", [(f"{letter}{r}", value_at(r, c))
                         for r in range(1, max_row + 1)]


def scan_worksheet(ws, *, min_group: int = 3) -> WorksheetReport:
    """Scan a worksheet for high-confidence flanked-singleton formula inconsistencies.

    Reads only (never writes, never materialises empty cells). Findings are deduplicated
    by cell and returned in deterministic top-to-bottom, left-to-right order, each carrying
    the axes that flagged it and — when the row and column scans disagree on the expected
    pattern — an explicit conflict with both expectations preserved.
    """
    report = WorksheetReport()
    normal_cache: dict[str, object] = {}
    analysed: set[str] = set()          # comparable formulas (pattern resolved, in a series)
    resolved_coords: set[str] = set()
    formula_coords: set[str] = set()
    merged: dict[str, dict] = {}        # coord -> {"observed": str, "by_axis": {axis: expected}}

    def normal_of(coord, value):
        nr = normal_cache.get(coord)
        if nr is None:
            nr = normalize_cell_value(value, coord)
            normal_cache[coord] = nr
        return nr

    for axis, line in _lines(ws):
        coords = [c for c, _v in line]
        values = [v for _c, v in line]

        if axis == "row":  # count coverage once (rows cover every cell exactly once)
            for coord, value in line:
                if _member_kind(value) == "formula":
                    formula_coords.add(coord)
                    nr = normal_of(coord, value)
                    if nr.coverage == RESOLVED:
                        report.resolved += 1
                        resolved_coords.add(coord)
                    else:
                        report.unsupported += 1
                        report.unsupported_cells.append((coord, nr.coverage))

        for start, end, kind in segment_line(values):
            if kind != "formula":
                continue
            seg_coords = coords[start:end + 1]
            seg_patterns = [normal_of(coords[i], values[i]).pattern for i in range(start, end + 1)]
            dom, count = dominant_pattern(seg_patterns)
            if dom is None or count < min_group:
                continue
            for coord, pattern in zip(seg_coords, seg_patterns):
                if pattern is not None:          # only genuinely comparable formulas
                    analysed.add(coord)
            for f in flanked_outliers(seg_coords, seg_patterns, dom):
                entry = merged.setdefault(f.coord, {"observed": f.observed, "by_axis": {}})
                entry["by_axis"][axis] = f.expected

    report.total_formulas = len(formula_coords)
    report.analysed_formulas = len(analysed)
    report.unsupported_cells.sort(key=lambda t: coordinate_to_tuple(t[0]))

    # Skipped = resolved formulas that never became comparable (isolated / short runs).
    for coord in sorted(resolved_coords - analysed, key=coordinate_to_tuple):
        report.skipped_cells.append((coord, "insufficient series context"))

    for coord, entry in merged.items():
        row, col = coordinate_to_tuple(coord)
        by_axis = entry["by_axis"]
        axes = tuple(sorted(by_axis))
        expected_by_axis = tuple((a, by_axis[a]) for a in axes)
        conflict = len(set(by_axis.values())) > 1
        where = " and ".join(axes)
        message = (f"Formula differs from the consistent pattern of its neighbours "
                   f"({where}). Potential inconsistency — review whether the difference "
                   f"is intentional.")
        if conflict:
            detail = "; ".join(f"{a} expects {e}" for a, e in expected_by_axis)
            message += (f" Note: the row and column scans disagree on the expected pattern "
                        f"({detail}); both are reported for review.")
        report.findings.append(WorksheetFinding(
            coord=coord, row=row, col=col, axes=axes,
            observed=entry["observed"], expected=expected_by_axis[0][1],
            expected_by_axis=expected_by_axis, conflict=conflict,
            confidence="high", message=message,
        ))
    report.findings.sort(key=lambda f: (f.row, f.col))
    return report
