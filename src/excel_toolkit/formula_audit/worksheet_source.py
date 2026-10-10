"""A small read-only worksheet interface so the detector is independent of its input source.

The detector only needs: sheet identity, bounds, and the value at each coordinate
(a formula string, a constant, or ``None``). Both a saved workbook (openpyxl) and a live
Excel workbook (a COM snapshot, built later) implement this one interface, so the *same*
analysis and confidence rules run against either — there is no second detection pathway.

    saved .xlsx ──openpyxl──┐
                            ├─▶ WorksheetSource ─▶ scan_worksheet (one detector)
    live Excel ──COM snap──┘

Reads are non-mutating by contract (see ``OpenpyxlWorksheet.value_at``).
"""

from __future__ import annotations

from typing import Protocol, runtime_checkable

from openpyxl.utils.cell import coordinate_to_tuple


@runtime_checkable
class WorksheetSource(Protocol):
    """Read-only view a detector needs. 1-based ``row``/``col`` like Excel/openpyxl."""

    @property
    def title(self) -> str: ...
    @property
    def max_row(self) -> int: ...
    @property
    def max_col(self) -> int: ...
    def value_at(self, row: int, col: int): ...  # formula str | constant | None


class OpenpyxlWorksheet:
    """Adapter over an openpyxl worksheet (a saved workbook)."""

    def __init__(self, ws):
        self._ws = ws

    @property
    def title(self) -> str:
        return self._ws.title

    @property
    def max_row(self) -> int:
        return self._ws.max_row

    @property
    def max_col(self) -> int:
        return self._ws.max_column

    def value_at(self, row: int, col: int):
        # ws._cells.get never creates a Cell (unlike ws.cell()/iter_rows) -> no mutation.
        cell = self._ws._cells.get((row, col))
        return cell.value if cell is not None else None


class InMemoryWorksheet:
    """A plain in-memory snapshot — the shape a live Excel COM read will supply.

    Holds exactly the formulas/constants present (including *unsaved* edits, which simply
    means the snapshot carries different values than the file on disk). Build from A1-keyed
    or (row, col)-keyed cells.
    """

    def __init__(self, cells: dict, title: str = "Sheet1"):
        self._cells: dict[tuple[int, int], object] = {}
        for key, value in cells.items():
            rc = key if isinstance(key, tuple) else coordinate_to_tuple(key)
            self._cells[rc] = value
        self._title = title
        self._max_row = max((r for r, _c in self._cells), default=0)
        self._max_col = max((c for _r, c in self._cells), default=0)

    @classmethod
    def from_openpyxl(cls, ws) -> "InMemoryWorksheet":
        cells = {rc: cell.value for rc, cell in ws._cells.items() if cell.value is not None}
        return cls(cells, title=ws.title)

    @property
    def title(self) -> str:
        return self._title

    @property
    def max_row(self) -> int:
        return self._max_row

    @property
    def max_col(self) -> int:
        return self._max_col

    def value_at(self, row: int, col: int):
        return self._cells.get((row, col))


def as_source(obj) -> WorksheetSource:
    """Return ``obj`` if it already implements the interface, else wrap an openpyxl ws."""
    if hasattr(obj, "value_at") and hasattr(obj, "max_row") and hasattr(obj, "max_col"):
        return obj
    return OpenpyxlWorksheet(obj)


def snapshot_from_com(com_ws) -> InMemoryWorksheet:
    """Build an InMemoryWorksheet from a live Excel COM worksheet object.

    This is the snapshot-conversion logic for the eventual Excel integration. It is tested
    only against a FAKE COM object offline; it is NOT run against real Excel here.

    Unresolved live-Excel assumptions (to confirm on the authorised test machine):
      * uses ``Range.Formula`` (always US A1 syntax with ',' separators, which the openpyxl
        tokenizer expects) — NOT ``FormulaLocal``;
      * ``Cell.HasFormula`` distinguishes formulas from constants;
      * ``UsedRange`` gives the bounds (``.Row``/``.Column`` top-left, ``.Rows.Count`` /
        ``.Columns.Count``); it can over-report and does not special-case merged cells or
        array/dynamic-array formulas (``Formula2`` / spill) — those remain future work;
      * 1-based indexing and ``Worksheet.Name`` for identity.
    """
    used = com_ws.UsedRange
    top, left = used.Row, used.Column
    n_rows, n_cols = used.Rows.Count, used.Columns.Count
    cells: dict[tuple[int, int], object] = {}
    for r in range(top, top + n_rows):
        for c in range(left, left + n_cols):
            cell = com_ws.Cells(r, c)
            if cell.HasFormula:
                cells[(r, c)] = cell.Formula
            else:
                value = cell.Value
                if value is not None:
                    cells[(r, c)] = value
    return InMemoryWorksheet(cells, title=com_ws.Name)
